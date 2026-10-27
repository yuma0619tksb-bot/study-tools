#!/usr/bin/env python3
"""授業資料の取り込み・索引・検索ツール。

使い方(リポジトリ直下で実行):
  python3 tools/materials.py ingest            # raw/ の新規・更新ファイルをテキスト化し、目次を更新
  python3 tools/materials.py status            # 要約待ち・OCR待ちの一覧
  python3 tools/materials.py search 語 [語..]   # 全資料をページ単位で検索(全語を含むページ)
  python3 tools/materials.py ocr [--subject S]  # OCR待ちページを画像化して一覧表示
  python3 tools/materials.py export-obsidian [--out DIR]   # 要約をObsidian用ノートに書き出し
  python3 tools/materials.py bundle [--subject S]          # claude.aiプロジェクトにアップする一式を作る

置き場所: materials/<科目>/raw/ に元ファイルを入れる。
生成物:   materials/<科目>/text/<名前>.md   ページ単位のテキスト(## p.N 見出し)
          materials/<科目>/digest/<名前>.md 要約+キーワード→ページ索引(Claudeが書く)
          materials/INDEX.md, materials/catalog.json
"""
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MAT = ROOT / "materials"
CATALOG = MAT / "catalog.json"
INDEX = MAT / "INDEX.md"
EXPORTS = ROOT / "exports"

OCR_MARK = "[OCR待ち]"
MIN_CHARS = 30  # これ未満の文字数のページはスキャンとみなす
PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".heic", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
DOCX_EXT = {".docx"}
PPTX_EXT = {".pptx"}
LEGACY_EXT = {".doc", ".ppt", ".odt", ".odp", ".rtf"}
SUPPORTED = PDF_EXT | IMG_EXT | DOCX_EXT | PPTX_EXT | LEGACY_EXT | {".txt", ".md"}


# ---------- 共通 ----------

def sha1(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load_catalog():
    if CATALOG.exists():
        return json.loads(CATALOG.read_text(encoding="utf-8"))
    return {}


def save_catalog(cat):
    CATALOG.write_text(json.dumps(cat, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def read_frontmatter(path):
    if not path.exists():
        return {}
    text = path.read_text(encoding="utf-8")
    m = re.match(r"---\n(.*?)\n---\n", text, re.S)
    if not m:
        return {}
    meta = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.strip()
    return meta


KIND_ORDER = {"講義資料": 0, "課題・演習": 1, "過去問・試験": 2}


def sort_key(e):
    return (KIND_ORDER.get(e["kind"], 9), e.get("lecture") or 999, e["stem"])


def guess_kind(name):
    n = name.lower()
    if re.search(r"過去問|試験|中間|期末|小テスト|exam|quiz|midterm|final", n):
        return "過去問・試験"
    if re.search(r"課題|レポート|宿題|演習|assignment|report|hw|homework", n):
        return "課題・演習"
    return "講義資料"


def guess_lecture_no(name):
    for pat in (r"第\s*(\d+)\s*回", r"(\d+)\s*回", r"(?:lec|lecture|week|wk|ch)[_\-\s]*(\d+)", r"^(\d{1,2})[_\-\s]"):
        m = re.search(pat, name, re.I)
        if m:
            return int(m.group(1))
    return None


def stem_for(raw_dir, path):
    # サブフォルダ内のファイルも衝突しないよう、相対パスを __ でつなぐ
    rel = path.relative_to(raw_dir).with_suffix("")
    return "__".join(rel.parts)


# ---------- 抽出 ----------

def extract_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for i in range(len(reader.pages)):
        # pdftotext -layout の方が日本語・表の崩れが少ない
        try:
            out = subprocess.run(
                ["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), str(path), "-"],
                capture_output=True, text=True, timeout=60,
            ).stdout
        except (FileNotFoundError, subprocess.TimeoutExpired):
            out = reader.pages[i].extract_text() or ""
        pages.append(out)
    return pages


def extract_docx(path):
    import docx
    d = docx.Document(str(path))
    # Wordにはページの概念がないので、見出しごと(なければ40段落ごと)に区切る
    sections, cur = [], []
    for p in d.paragraphs:
        if p.style.name.lower().startswith(("heading", "見出し")) and cur:
            sections.append("\n".join(cur))
            cur = []
        if p.text.strip():
            cur.append(p.text)
        if len(cur) >= 40:
            sections.append("\n".join(cur))
            cur = []
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        cur.append("\n".join(rows))
    if cur:
        sections.append("\n".join(cur))
    return sections or [""]


def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    pages = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for r in shape.table.rows:
                    texts.append(" | ".join(c.text for c in r.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("(ノート) " + slide.notes_slide.notes_text_frame.text)
        pages.append("\n".join(texts))
    return pages


def convert_legacy(path):
    """doc/ppt などを LibreOffice で PDF に変換して抽出する。"""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
                       capture_output=True, timeout=180)
        pdf = Path(tmp) / (path.stem + ".pdf")
        if not pdf.exists():
            raise RuntimeError("LibreOfficeでの変換に失敗")
        return extract_pdf(pdf)


def extract(path):
    ext = path.suffix.lower()
    if ext in PDF_EXT:
        return extract_pdf(path)
    if ext in DOCX_EXT:
        return extract_docx(path)
    if ext in PPTX_EXT:
        return extract_pptx(path)
    if ext in LEGACY_EXT:
        return convert_legacy(path)
    if ext in IMG_EXT:
        return [""]
    return [path.read_text(encoding="utf-8", errors="replace")]


def normalize(text):
    text = text.replace("\x0c", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------- ingest ----------

def write_text_file(subject, path, stem, pages, digest):
    out = MAT / subject / "text" / f"{stem}.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    unit = "スライド" if path.suffix.lower() in PPTX_EXT else ("節" if path.suffix.lower() in DOCX_EXT else "p")
    ocr = []
    lines = [
        "---",
        f"source: {path.relative_to(MAT / subject).as_posix()}",
        f"pages: {len(pages)}",
        f"sha1: {digest}",
        "---",
        f"# {path.name}",
        "",
    ]
    for i, raw in enumerate(pages, 1):
        body = normalize(raw)
        lines.append(f"## p.{i}")
        scannable = path.suffix.lower() in PDF_EXT | IMG_EXT | LEGACY_EXT
        if scannable and len(re.sub(r"\s", "", body)) < MIN_CHARS:
            ocr.append(i)
            lines.append(f"{OCR_MARK} {unit}{i}は文字が取れませんでした(スキャン/画像)。`materials.py ocr` で画像化して読み取ってください。")
            if body:
                lines.append(body)
        else:
            lines.append(body)
        lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")
    return out, ocr


def count_ocr_pending(text_path):
    if not text_path.exists():
        return []
    pending, page = [], None
    for line in text_path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"## p\.(\d+)", line)
        if m:
            page = int(m.group(1))
        elif line.startswith(OCR_MARK) and page:
            pending.append(page)
    return pending


def subjects():
    if not MAT.exists():
        return []
    return sorted(p.name for p in MAT.iterdir() if p.is_dir() and (p / "raw").is_dir())


def cmd_ingest(args):
    cat = load_catalog()
    seen = set()
    changed = 0
    for subject in subjects():
        if args.subject and subject != args.subject:
            continue
        raw_dir = MAT / subject / "raw"
        for path in sorted(raw_dir.rglob("*")):
            if not path.is_file() or path.name.startswith(".") or path.suffix.lower() not in SUPPORTED:
                continue
            stem = stem_for(raw_dir, path)
            key = f"{subject}/{stem}"
            seen.add(key)
            digest = sha1(path)
            text_path = MAT / subject / "text" / f"{stem}.md"
            entry = cat.get(key, {})
            # 内容が同じなら再抽出しない(手で書き込んだOCR結果を守るため)
            if not args.force and entry.get("sha1") == digest and text_path.exists():
                entry["ocr_pending"] = count_ocr_pending(text_path)
                cat[key] = entry
                continue
            try:
                pages = extract(path)
            except Exception as e:  # 1ファイルの失敗で全体を止めない
                print(f"  ✗ {key}: {e}", file=sys.stderr)
                continue
            _, ocr = write_text_file(subject, path, stem, pages, digest)
            entry.update({
                "subject": subject,
                "stem": stem,
                "source": path.relative_to(MAT / subject).as_posix(),
                "sha1": digest,
                "pages": len(pages),
                "kind": entry.get("kind_manual") or guess_kind(path.name),
                "lecture": entry.get("lecture_manual") or guess_lecture_no(path.stem),
                "ingested": date.today().isoformat(),
                "ocr_pending": ocr,
            })
            cat[key] = entry
            changed += 1
            note = f"(OCR待ち {len(ocr)}p)" if ocr else ""
            print(f"  ✓ {key}  {len(pages)}p {note}")
    if not args.subject:
        for key in [k for k in cat if k not in seen]:
            print(f"  - {key}: 元ファイルが無いので目録から削除")
            del cat[key]
    refresh_digest_state(cat)
    save_catalog(cat)
    write_index(cat)
    print(f"取り込み {changed} 件 / 合計 {len(cat)} 件。materials/INDEX.md を更新しました。")


def refresh_digest_state(cat):
    for key, e in cat.items():
        d = MAT / e["subject"] / "digest" / f"{e['stem']}.md"
        meta = read_frontmatter(d)
        if not meta:
            e["digest"] = "未作成"
            e.pop("summary", None)
        else:
            e["digest"] = "最新" if meta.get("sha1") == e["sha1"] else "古い(元ファイル更新)"
            e["summary"] = meta.get("summary", "")


def write_index(cat):
    lines = [
        "# 資料目次",
        "",
        "Claudeはまずこのファイルを読み、必要な資料の digest → text の該当ページ の順に絞って読む。",
        "更新: `python3 tools/materials.py ingest`(手で編集しない)",
        "",
    ]
    by_subject = {}
    for e in cat.values():
        by_subject.setdefault(e["subject"], []).append(e)
    for subject in sorted(by_subject):
        items = sorted(by_subject[subject], key=sort_key)
        lines += [f"## {subject}", "", "| 回 | 種類 | 資料 | ページ | 要約 | 一言 |", "|---|---|---|---|---|---|"]
        for e in items:
            lec = e.get("lecture") or "-"
            pages = f"{e['pages']}" + (f" (OCR待ち{len(e['ocr_pending'])})" if e.get("ocr_pending") else "")
            link = f"[{e['stem']}]({e['subject']}/text/{e['stem']}.md)"
            dg = e["digest"]
            if dg != "未作成":
                dg = f"[{dg}]({e['subject']}/digest/{e['stem']}.md)"
            summary = (e.get("summary") or "").replace("|", "/")
            lines.append(f"| {lec} | {e['kind']} | {link} | {pages} | {dg} | {summary} |")
        lines.append("")
    if not by_subject:
        lines.append("(まだ資料がありません。materials/<科目>/raw/ にファイルを入れて ingest してください)")
    INDEX.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------- status / search / ocr ----------

def cmd_status(args):
    cat = load_catalog()
    refresh_digest_state(cat)
    todo_digest = [k for k, e in cat.items() if e["digest"] != "最新" and (not args.subject or e["subject"] == args.subject)]
    todo_ocr = [(k, e["ocr_pending"]) for k, e in cat.items() if e.get("ocr_pending") and (not args.subject or e["subject"] == args.subject)]
    if args.quiet and not todo_digest and not todo_ocr:
        return
    print(f"資料 {len(cat)} 件")
    print(f"\n■ 要約(digest)が未作成・古い: {len(todo_digest)} 件")
    for k in todo_digest:
        print(f"  {k}  [{cat[k]['digest']}]")
    print(f"\n■ OCR待ち: {len(todo_ocr)} 件")
    for k, pages in todo_ocr:
        print(f"  {k}  p.{','.join(map(str, pages))}")


def iter_pages(text_path):
    page, buf = None, []
    for line in text_path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"## p\.(\d+)", line)
        if m:
            if page is not None:
                yield page, "\n".join(buf)
            page, buf = int(m.group(1)), []
        elif page is not None:
            buf.append(line)
    if page is not None:
        yield page, "\n".join(buf)


def cmd_search(args):
    terms = [t.lower() for t in args.terms]
    hits = 0
    for subject in subjects():
        if args.subject and subject != args.subject:
            continue
        for tp in sorted((MAT / subject / "text").glob("*.md")):
            for page, body in iter_pages(tp):
                low = body.lower()
                if all(t in low for t in terms):
                    hits += 1
                    i = low.find(terms[0])
                    snip = body[max(0, i - 40): i + 80].replace("\n", " ")
                    print(f"{subject}/{tp.stem}  p.{page}: …{snip}…")
                    if hits >= args.limit:
                        print(f"(上限 {args.limit} 件で打ち切り。--limit で変更)")
                        return
    if hits == 0:
        print("該当なし。言い換え(英語・略語・表記ゆれ)でも検索してください。")


def cmd_ocr(args):
    """OCR待ちページを PNG にして場所を表示する。読み取りは Claude(画像を読む)が行う。"""
    cat = load_catalog()
    for key, e in sorted(cat.items()):
        if not e.get("ocr_pending") or (args.subject and e["subject"] != args.subject):
            continue
        src = MAT / e["subject"] / e["source"]
        outdir = MAT / e["subject"] / "pages" / e["stem"]
        outdir.mkdir(parents=True, exist_ok=True)
        ext = src.suffix.lower()
        for p in e["ocr_pending"]:
            png = outdir / f"p{p:03d}.png"
            if png.exists():
                pass
            elif ext in IMG_EXT:
                shutil.copy(src, outdir / f"p{p:03d}{ext}")
                png = outdir / f"p{p:03d}{ext}"
            elif ext in PDF_EXT:
                subprocess.run(["pdftoppm", "-r", str(args.dpi), "-png", "-singlefile",
                                "-f", str(p), "-l", str(p), str(src), str(png.with_suffix(""))], check=False)
            else:
                continue
            print(f"{key} p.{p}: {png.relative_to(ROOT)}")
    print("\n読み取ったら text/<名前>.md の該当 `## p.N` の [OCR待ち] 行を本文に置き換え、ingest を再実行。")


# ---------- export ----------

def cmd_export_obsidian(args):
    cat = load_catalog()
    refresh_digest_state(cat)
    out = Path(args.out) if args.out else EXPORTS / "obsidian"
    n = 0
    by_subject = {}
    for e in cat.values():
        if args.subject and e["subject"] != args.subject:
            continue
        by_subject.setdefault(e["subject"], []).append(e)
    for subject, items in by_subject.items():
        base = out / subject / "資料"
        base.mkdir(parents=True, exist_ok=True)
        rows = []
        for e in sorted(items, key=sort_key):
            d = MAT / subject / "digest" / f"{e['stem']}.md"
            title = e["stem"]
            if d.exists():
                body = re.sub(r"^---\n.*?\n---\n", "", d.read_text(encoding="utf-8"), flags=re.S)
                note = [
                    "---",
                    f"科目: {subject}",
                    f"種類: {e['kind']}",
                    f"回: {e.get('lecture') or ''}",
                    f"元ファイル: {e['source']}",
                    f"tags: [資料, {subject}]",
                    "---",
                    body.strip(),
                    "",
                ]
                (base / f"{title}.md").write_text("\n".join(note), encoding="utf-8")
                n += 1
            rows.append(f"- [[{title}]] — {e['kind']} {('第' + str(e['lecture']) + '回') if e.get('lecture') else ''} {e.get('summary') or '(要約未作成)'}")
        (out / subject / f"{subject} 資料一覧.md").write_text(
            f"# {subject} 資料一覧\n\n" + "\n".join(rows) + "\n", encoding="utf-8")
    print(f"Obsidian用ノート {n} 件を {out} に書き出しました。")


def cmd_bundle(args):
    """claude.aiプロジェクトの「ナレッジ」にアップする一式。元PDFより軽く、検索しやすい。"""
    cat = load_catalog()
    refresh_digest_state(cat)
    for subject in subjects():
        if args.subject and subject != args.subject:
            continue
        out = EXPORTS / "claude_project" / subject
        if out.exists():
            shutil.rmtree(out)
        out.mkdir(parents=True)
        items = [e for e in cat.values() if e["subject"] == subject]
        idx = [f"# {subject} 資料目次", "", "| 回 | 種類 | 資料 | ページ | 一言 |", "|---|---|---|---|---|"]
        for e in sorted(items, key=sort_key):
            idx.append(f"| {e.get('lecture') or '-'} | {e['kind']} | {e['stem']} | {e['pages']} | {e.get('summary') or ''} |")
            parts = []
            d = MAT / subject / "digest" / f"{e['stem']}.md"
            if d.exists():
                parts.append(d.read_text(encoding="utf-8"))
            parts.append((MAT / subject / "text" / f"{e['stem']}.md").read_text(encoding="utf-8"))
            (out / f"{e['stem']}.md").write_text("\n\n".join(parts), encoding="utf-8")
        (out / "00_資料目次.md").write_text("\n".join(idx) + "\n", encoding="utf-8")
        print(f"{subject}: {len(items)} 件 → {out.relative_to(ROOT)}/")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ingest"); p.add_argument("--subject"); p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_ingest)
    p = sub.add_parser("status"); p.add_argument("--subject")
    p.add_argument("--quiet", action="store_true", help="未処理が無ければ何も出さない")
    p.set_defaults(func=cmd_status)
    p = sub.add_parser("search"); p.add_argument("terms", nargs="+"); p.add_argument("--subject")
    p.add_argument("--limit", type=int, default=30); p.set_defaults(func=cmd_search)
    p = sub.add_parser("ocr"); p.add_argument("--subject"); p.add_argument("--dpi", type=int, default=150)
    p.set_defaults(func=cmd_ocr)
    p = sub.add_parser("export-obsidian"); p.add_argument("--subject"); p.add_argument("--out")
    p.set_defaults(func=cmd_export_obsidian)
    p = sub.add_parser("bundle"); p.add_argument("--subject"); p.set_defaults(func=cmd_bundle)
    args = ap.parse_args()
    MAT.mkdir(exist_ok=True)
    args.func(args)


if __name__ == "__main__":
    main()
